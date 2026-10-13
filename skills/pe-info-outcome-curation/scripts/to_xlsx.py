#!/usr/bin/env python3
"""Combine one paper's two PE deliverables into a single .xlsx (optional).

Reads <dir>/pe-info-outcome-curation.study_info.csv and
<dir>/pe-info-outcome-curation.csv and writes
<dir>/pe-info-outcome-curation.xlsx with two sheets, "Study info" and "Outcomes".
Every cell is written as text, so values such as ".67" or "<0.001" stay exactly
as extracted.

Usage:
    python to_xlsx.py <output-dir>/<pmid>

Needs openpyxl (in the repo's environment.yml, or `pip install openpyxl`).
"""
import csv
import os
import sys

SKILL = "pe-info-outcome-curation"


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 1:
        print(__doc__, file=sys.stderr)
        return 2
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("error: openpyxl is not installed (pip install openpyxl)", file=sys.stderr)
        return 2

    d = argv[0]
    sheets = [("Study info", os.path.join(d, f"{SKILL}.study_info.csv")),
              ("Outcomes", os.path.join(d, f"{SKILL}.csv"))]
    wb = Workbook()
    wb.remove(wb.active)
    for name, path in sheets:
        if not os.path.exists(path):
            print(f"error: missing {path}", file=sys.stderr)
            return 2
        ws = wb.create_sheet(name)
        with open(path, encoding="utf-8", newline="") as fh:
            for i, row in enumerate(csv.reader(fh)):
                ws.append(row)
                for cell in ws[ws.max_row]:
                    cell.number_format = "@"
                    if i == 0:
                        cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    out = os.path.join(d, f"{SKILL}.xlsx")
    wb.save(out)
    print(f"wrote {os.path.abspath(out)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
